import shutil

import openpyxl
import pytest

from src.parser import ContainerLine, ShipmentRecord
from src.tracker_sheet import TrackerError, append_shipment

SHEET_NAME = "Incoming vessels-2026"


def _make_workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(SHEET_NAME)
    ws.append(["title"])
    ws.append(["subtitle"])
    ws.append([])
    ws.append(
        ["S.NO.", "VESSEL NAME", None, "ETA", "ARRIVAL DATE", "Unit", "UNIT COUNT",
         "QTY(MT)", "OFFLOAD QTY", "OFFLOAD DATE", "ORDER #", "DN#", "ISOT#", "REMARKS"]
    )
    for i in range(1, 21):
        ws.append([i, None, None, None, None, "iso tanks", None, None, None, None, None, None, None, None])
    wb.save(path)


def _sample_shipment():
    return ShipmentRecord(
        order_number="57731272",
        vessel_name="CORNELIA MAERSK",
        containers=[
            ContainerLine(container_no="HGTU2418040", delivery_no="93065721", net_wt_kg=21736.0),
            ContainerLine(container_no="RLTU2626928", delivery_no="93065722", net_wt_kg=21745.0),
        ],
    )


def test_appends_into_next_blank_template_rows(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _make_workbook(path)

    rows = append_shipment(str(path), SHEET_NAME, _sample_shipment())
    assert rows == [5, 6]

    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_NAME]
    assert ws.cell(row=5, column=2).value == "CORNELIA MAERSK"
    assert ws.cell(row=5, column=7).value == 2
    assert ws.cell(row=6, column=2).value is None  # second row: no vessel/unit count
    assert ws.cell(row=6, column=12).value == 93065722


def test_second_write_lands_in_the_next_free_rows(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _make_workbook(path)
    append_shipment(str(path), SHEET_NAME, _sample_shipment())

    next_shipment = ShipmentRecord(
        order_number="57731280",
        vessel_name="MSC EXAMPLE",
        containers=[ContainerLine(container_no="TCLU1111111", delivery_no="93065999", net_wt_kg=20000.0)],
    )
    rows = append_shipment(str(path), SHEET_NAME, next_shipment)
    assert rows == [7]


def test_refuses_to_duplicate_an_already_logged_delivery_number(tmp_path):
    path = tmp_path / "tracker.xlsx"
    _make_workbook(path)
    append_shipment(str(path), SHEET_NAME, _sample_shipment())

    with pytest.raises(TrackerError, match="already"):
        append_shipment(str(path), SHEET_NAME, _sample_shipment())
